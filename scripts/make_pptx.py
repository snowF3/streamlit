"""
동네 엑스레이 발표자료 — 도형 기반 고도화 버전
"""
from pptx import Presentation
from pptx.util import Pt, Inches, Cm, Emu
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pathlib import Path

BASE = Path(__file__).resolve().parent.parent
TEMPLATE = BASE / "Snowflake.pptx"
OUTPUT = BASE / "동네엑스레이_발표자료.pptx"
ICONS = Path(r"C:\Users\bbba2\OneDrive\바탕 화면\토스_포트폴리오\icons")
prs = Presentation(str(TEMPLATE))

PURPLE = RGBColor.from_string("6366F1")
GRAY = RGBColor.from_string("666666")
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
BLACK = RGBColor(0x33, 0x33, 0x33)


def _set_text(shape, lines):
    tf = shape.text_frame
    tf.word_wrap = True
    for p in tf.paragraphs:
        p.text = ""
    for i, (text, size, bold, color) in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        run = p.add_run()
        run.text = text
        run.font.size = Pt(size)
        run.font.bold = bold
        if color:
            run.font.color.rgb = RGBColor.from_string(color) if isinstance(color, str) else color


def _find(slide, keyword):
    for s in slide.shapes:
        if s.has_text_frame and keyword in s.text_frame.text:
            return s
    return None


def _find_box_before(slide, label_keyword):
    shapes = list(slide.shapes)
    for i, s in enumerate(shapes):
        if s.has_text_frame and label_keyword in s.text_frame.text:
            for j in range(i-1, -1, -1):
                sh = shapes[j]
                if sh.shape_type == 1 and sh.has_text_frame and sh.width > 1000000:
                    return sh
    return None


def _add_rounded_label(slide, left, top, width, height, text, font_size=10):
    """보라색 라운드 네모 라벨 추가"""
    shape = slide.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE, left, top, width, height
    )
    shape.fill.solid()
    shape.fill.fore_color.rgb = PURPLE
    shape.line.fill.background()
    # 라운드 조절
    if hasattr(shape, 'adjustments') and len(shape.adjustments) > 0:
        shape.adjustments[0] = 0.3
    tf = shape.text_frame
    tf.word_wrap = True
    tf.paragraphs[0].alignment = PP_ALIGN.CENTER
    run = tf.paragraphs[0].add_run()
    run.text = text
    run.font.size = Pt(font_size)
    run.font.bold = True
    run.font.color.rgb = WHITE
    return shape


def _add_pipe_box(slide, left, top, width, height, title, desc, fill_color=None):
    """파이프라인 단계 박스"""
    shape = slide.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE, left, top, width, height
    )
    if fill_color:
        shape.fill.solid()
        shape.fill.fore_color.rgb = fill_color
    else:
        shape.fill.solid()
        shape.fill.fore_color.rgb = RGBColor(0xF0, 0xF0, 0xF5)
    shape.line.color.rgb = RGBColor(0xDD, 0xDD, 0xDD)
    shape.line.width = Pt(0.5)

    tf = shape.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    run = p.add_run()
    run.text = title
    run.font.size = Pt(8)
    run.font.bold = True
    run.font.color.rgb = BLACK if not fill_color else WHITE
    if desc:
        p2 = tf.add_paragraph()
        p2.alignment = PP_ALIGN.CENTER
        run2 = p2.add_run()
        run2.text = desc
        run2.font.size = Pt(6)
        run2.font.color.rgb = GRAY if not fill_color else RGBColor(0xDD, 0xDD, 0xDD)
    return shape


def _add_arrow(slide, left, top, width, height):
    """아래 화살표"""
    shape = slide.shapes.add_shape(
        MSO_SHAPE.DOWN_ARROW, left, top, width, height
    )
    shape.fill.solid()
    shape.fill.fore_color.rgb = PURPLE
    shape.line.fill.background()
    return shape


# ═══════════════════════════════════════
# 슬라이드 1: 표지
# ═══════════════════════════════════════
s1 = prs.slides[0]
sh = _find(s1, "Name of Individual")
if sh: _set_text(sh, [("동네 엑스레이", 36, True, None)])
sh = _find(s1, "Organization Name")
if sh: _set_text(sh, [("XR-AI  |  오용석 · 박성우 · 박재성", 14, False, None)])
sh = _find(s1, "Apr 2026")
if sh: _set_text(sh, [("AI 기반 상권 분석 플랫폼  |  April 2026", 11, False, None)])

# ═══════════════════════════════════════
# 슬라이드 4: 과제 정의 — 도형 기반
# ═══════════════════════════════════════
s4 = prs.slides[3]

# 기존 라벨 텍스트 교체 (보라색 라운드 배경은 별도 도형으로)
labels = {
    "문제 정의": ("문제 정의", Cm(0.7)),
    "가설": ("가설", Cm(6.9)),
    "아키텍처 구조": ("아키텍처", Cm(13.1)),
    "사용된 기술 스택": ("기술 스택", Cm(19.3)),
}

# 라벨 위에 보라색 라운드 네모 추가
for label_text, (new_text, left_pos) in labels.items():
    sh = _find(s4, label_text)
    if sh:
        # 기존 라벨 텍스트 비우기
        _set_text(sh, [("", 1, False, None)])
    # 라운드 라벨 추가
    _add_rounded_label(s4, left_pos, Cm(1.4), Cm(5.0), Cm(0.8), new_text, 9)

# 기존 내용 박스 채우기
card_map = {
    "Problem Statement:": [
        ("", 4, False, None),
        ("자영업자 10명 중 7명이 감에 의존하여", 8, False, None),
        ("상권 입지를 선정하고 있습니다.", 8, False, None),
        ("", 4, False, None),
        ("기존 상권 분석 서비스는 과거 데이터만 제공하여", 8, False, None),
        ("실시간 트렌드와 미래 예측이 불가능합니다.", 8, False, None),
        ("", 4, False, None),
        ("유동인구·소비·부동산 데이터는 존재하나", 8, False, None),
        ("이를 통합 분석하는 도구가 부재합니다.", 8, False, None),
        ("", 6, False, None),
        ("Theme", 8, True, "6366F1"),
        ("상권 분석 / 인구 이동 예측", 8, False, "666666"),
    ],
    "State your hypothesis": [
        ("", 4, False, None),
        ("5개 선행지표를 종합하면", 8, False, None),
        ("'다음에 뜰 동네'를 예측할 수 있다", 8, True, None),
        ("", 4, False, None),
        ("방문인구 25%", 8, False, None),
        ("카페·식음료 매출 20%", 8, False, None),
        ("유동인구 총합 20%", 8, False, None),
        ("부동산 매매가 20%", 8, False, None),
        ("인터넷 신규설치 15%", 8, False, None),
        ("", 4, False, None),
        ("118개 동네 x 59개월 = 6,962건", 8, True, "6366F1"),
    ],
}
for shape in s4.shapes:
    if shape.has_text_frame:
        for key, lines in card_map.items():
            if key in shape.text_frame.text:
                _set_text(shape, lines)
                break

# 아키텍처: 기존 텍스트 비우고 파이프라인 도형 추가
arch_shape = _find(s4, "Insert Snowflake Architecture")
if arch_shape:
    _set_text(arch_shape, [("", 1, False, None)])  # 비우기

# 파이프라인 도형들
arch_left = Cm(13.3)
box_w = Cm(5.0)
box_h = Cm(1.5)
arrow_w = Cm(0.5)
arrow_h = Cm(0.5)

y = Cm(2.8)
_add_pipe_box(s4, arch_left, y, box_w, box_h, "Snowflake Marketplace", "SPH · 리치고 · 아정당 (5개)", PURPLE)
y += box_h + Cm(0.1)
_add_arrow(s4, arch_left + box_w/2 - arrow_w/2, y, arrow_w, arrow_h)
y += arrow_h + Cm(0.1)
_add_pipe_box(s4, arch_left, y, box_w, box_h, "Snowflake SQL", "전처리 + 핫플 점수 실시간 계산")
y += box_h + Cm(0.1)
_add_arrow(s4, arch_left + box_w/2 - arrow_w/2, y, arrow_w, arrow_h)
y += arrow_h + Cm(0.1)
_add_pipe_box(s4, arch_left, y, box_w, box_h, "Cortex AI (GPT-5-mini)", "자연어 → SQL → 인사이트")
y += box_h + Cm(0.1)
_add_arrow(s4, arch_left + box_w/2 - arrow_w/2, y, arrow_w, arrow_h)
y += arrow_h + Cm(0.1)
_add_pipe_box(s4, arch_left, y, box_w, box_h, "Streamlit in Snowflake", "피드 / AI에이전트 / 디지털트윈", PURPLE)

# 기술 스택: 아이콘 + 텍스트
tech_shape = _find(s4, "Add logos or names")
if tech_shape:
    _set_text(tech_shape, [("", 1, False, None)])  # 비우기

tech_left = Cm(19.5)
tech_items = [
    ("Snowflake Marketplace", None),
    ("Cortex AI (GPT-5-mini)", None),
    ("Streamlit", "streamlit.png"),
    ("Python 3.11", "python.png"),
    ("PyDeck 3D / Plotly", None),
]

ty = Cm(3.0)
icon_size = Cm(0.8)
for name, icon_file in tech_items:
    # 아이콘
    if icon_file and (ICONS / icon_file).exists():
        s4.shapes.add_picture(str(ICONS / icon_file), tech_left, ty, icon_size, icon_size)
        # 텍스트
        txt_box = s4.shapes.add_textbox(tech_left + Cm(1.0), ty, Cm(4.0), icon_size)
        tf = txt_box.text_frame
        tf.vertical_anchor = MSO_ANCHOR.MIDDLE
        run = tf.paragraphs[0].add_run()
        run.text = name
        run.font.size = Pt(8)
        run.font.bold = True
    else:
        # 아이콘 없으면 보라색 원 + 텍스트
        dot = s4.shapes.add_shape(MSO_SHAPE.OVAL, tech_left, ty + Cm(0.15), Cm(0.5), Cm(0.5))
        dot.fill.solid()
        dot.fill.fore_color.rgb = PURPLE
        dot.line.fill.background()
        txt_box = s4.shapes.add_textbox(tech_left + Cm(0.7), ty, Cm(4.3), icon_size)
        tf = txt_box.text_frame
        tf.vertical_anchor = MSO_ANCHOR.MIDDLE
        run = tf.paragraphs[0].add_run()
        run.text = name
        run.font.size = Pt(8)
        run.font.bold = True
    ty += Cm(1.2)

# "100% Snowflake Native" 뱃지
_add_rounded_label(s4, tech_left, ty + Cm(0.3), Cm(4.5), Cm(0.7), "100% Snowflake Native", 8)

# ═══════════════════════════════════════
# 인사이트 슬라이드 (5, 6, 7)
# ═══════════════════════════════════════
insights = [
    {
        "title": "인사이트 피드 — 다음에 뜰 동네 예측",
        "bullets": [
            ("5개 선행지표(방문인구 25% · 카페매출 20% · 유동인구 20% · 매매가 20% · 신규설치 15%)를 가중합하여 핫플 점수 산출", 8, False, None),
            ("월별 상위/하위 동네를 자동 탐지하고, '왜 올랐을까?' 변동 원인을 구체적 수치(유동인구 -5.3%, 카드매출 +92.8%)와 함께 자동 분석", 8, False, None),
            ("월별 누적 점수 추이 차트로 장기 상권 트렌드 파악 | 서울 3개구 118개 법정동 x 59개월 = 6,962건 Snowflake 실시간 계산", 8, False, None),
        ],
        "effect": [
            ("자영업자 입지 선정 혁신", 9, True, None),
            ("감이 아닌 데이터 근거로 '어디에 가게를 열까' 판단", 8, False, "666666"),
            ("잘못된 입지 선정으로 인한 폐업률 감소 기대", 8, False, "666666"),
            ("", 4, False, None),
            ("상권 변화 조기 포착", 9, True, None),
            ("월별 시그널 추적으로 상권 성장/쇠퇴를 6~12개월 전 감지", 8, False, "666666"),
            ("기존 임대계약 갱신·철수 의사결정에 선제적 대응", 8, False, "666666"),
            ("", 4, False, None),
            ("투자·정책 의사결정 지원", 9, True, None),
            ("부동산 투자자: 상권 성장 시그널 기반 투자 타이밍 포착", 8, False, "666666"),
            ("지자체: 상권 활성화 정책 수립 시 객관적 근거 자료 활용", 8, False, "666666"),
        ],
        "plan": [
            ("Phase 1: 데이터 확장", 9, True, None),
            ("서울 전역 → 수도권 → 전국 35,000개+ 법정동", 8, False, "666666"),
            ("Snowflake Marketplace 추가 데이터셋 연동", 8, False, "666666"),
            ("", 4, False, None),
            ("Phase 2: 예측 모델 고도화", 9, True, None),
            ("전월대비 변동률 → Snowpark ML 시계열 예측", 8, False, "666666"),
            ("3~6개월 후 핫플 점수 예측", 8, False, "666666"),
            ("", 4, False, None),
            ("Phase 3: 실시간 데이터 연동", 9, True, None),
            ("배달앱·SNS 주문량/언급량 추가", 8, False, "666666"),
            ("실시간 상권 트렌드 반영", 8, False, "666666"),
        ],
    },
    {
        "title": "AI 에이전트 — 자연어로 상권 분석",
        "bullets": [
            ("Cortex AI(GPT-5-mini)로 자연어 질문을 자동 SQL 변환 → Snowflake에서 직접 실행 → 즉시 결과 제공", 8, False, None),
            ("의도 분류(lookup/compare/trend/simulate) 기반 맥락 분석 자동 수행 + 대화 히스토리로 후속 질문 연속 처리", 8, False, None),
            ('"중구에서 카페 매출 가장 높은 동네는?" / "서초구 반포동 유동인구 추이 알려줘" / "영등포구 핫플 점수 비교해줘"', 8, False, "666666"),
        ],
        "effect": [
            ("데이터 분석 민주화", 9, True, None),
            ("SQL·코딩 없이 자연어로 누구나 상권 데이터 분석", 8, False, "666666"),
            ("데이터 분석가 없는 소상공인도 직접 인사이트 도출", 8, False, "666666"),
            ("", 4, False, None),
            ("의사결정 속도 99% 단축", 9, True, None),
            ("기존: 수집→정제→분석→보고 (수 시간~수 일)", 8, False, "666666"),
            ("AI 에이전트: 질문→답변 (수 초)", 8, False, "666666"),
            ("", 4, False, None),
            ("맥락 기반 심층 분석", 9, True, None),
            ("대화형으로 점점 깊은 분석 수행", 8, False, "666666"),
            ("이전 답변 맥락 유지하여 연속적 의사결정", 8, False, "666666"),
        ],
        "plan": [
            ("시뮬레이션 질의 지원", 9, True, None),
            ('"카페 열면 예상 월 매출은?" 가상 시나리오', 8, False, "666666"),
            ("핫플 점수 기반 매출 예측 모델 연동", 8, False, "666666"),
            ("", 4, False, None),
            ("멀티모달 출력 강화", 9, True, None),
            ("차트·지도·비교 테이블 자동 생성", 8, False, "666666"),
            ("Streamlit 위젯과 인터랙티브 연동", 8, False, "666666"),
            ("", 4, False, None),
            ("B2B 서비스화", 9, True, None),
            ("프랜차이즈 본사: 출점 후보지 자동 추천 API", 8, False, "666666"),
            ("부동산 중개업체: 상권 분석 리포트 자동 생성", 8, False, "666666"),
        ],
    },
    {
        "title": "디지털 트윈 — 살아있는 서울",
        "bullets": [
            ("실제 유동인구·소득·소비 데이터 기반 합성 시민을 생성하고 시간대별 거주지→직장→상권 이동을 PyDeck 3D로 시뮬레이션", 8, False, None),
            ("6가지 라이프스타일(싱글/신혼/영유아/청소년/성인자녀/실버) 기반 행동 패턴 반영 — 유형별 소비·이동·활동 시간대 차별화", 8, False, None),
            ('"오전 11시에 서초동에 어떤 사람들이 모이는가?" / "퇴근 시간대 중구 상권의 주요 고객층은?" — 직관적 시각 확인', 8, False, "666666"),
        ],
        "effect": [
            ("시간대별 타겟 고객 파악", 9, True, None),
            ("점심: 직장인 → 저녁: 거주민 등 시간대별 프로파일", 8, False, "666666"),
            ("영업시간·메뉴·마케팅을 고객 특성에 맞춰 최적화", 8, False, "666666"),
            ("", 4, False, None),
            ("입지 리스크 사전 검증", 9, True, None),
            ("오픈 전 시뮬레이션으로 예상 고객 수·유형 확인", 8, False, "666666"),
            ("잘못된 입지로 인한 초기 투자 손실 방지", 8, False, "666666"),
            ("", 4, False, None),
            ("도시 계획·마케팅 활용", 9, True, None),
            ("지자체: 보행자 동선 → 도로·시설 배치 최적화", 8, False, "666666"),
            ("마케팅: 옥외 광고 위치·시간대 최적화", 8, False, "666666"),
        ],
        "plan": [
            ("What-if 시뮬레이션", 9, True, None),
            ('"지하철역 신설 시 유동인구 변화 예측"', 8, False, "666666"),
            ('"대형 오피스 입주 시 상권 변화"', 8, False, "666666"),
            ("", 4, False, None),
            ("실시간 GPS 연동", 9, True, None),
            ("통신사 위치 데이터로 실제 이동 반영", 8, False, "666666"),
            ("합성→실제 시민 데이터로 정확도 향상", 8, False, "666666"),
            ("", 4, False, None),
            ("확장현실(XR) 연계", 9, True, None),
            ("VR: 가상 공간에서 상권 입지 체험", 8, False, "666666"),
            ("AR: 현장에서 실시간 상권 데이터 오버레이", 8, False, "666666"),
        ],
    },
]

for idx, data in enumerate(insights):
    slide = prs.slides[4 + idx]
    sh = _find(slide, "Insight and Recommendations")
    if sh: _set_text(sh, [(data["title"], 22, True, None)])
    sh = _find(slide, "Maximum 3 insights")
    if sh: _set_text(sh, data["bullets"])
    box = _find_box_before(slide, "기대 효과")
    if box: _set_text(box, data["effect"])
    box = _find_box_before(slide, "향후 실행 방안")
    if box: _set_text(box, data["plan"])

# 슬라이드 8: Thank You (원본 유지)

prs.save(str(OUTPUT))
print(f"Done: {OUTPUT}")
